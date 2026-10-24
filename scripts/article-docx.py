#!/usr/bin/env python3
"""Build a .docx of docs/小黑盒文章.md (alternative import for the 小黑盒 editor), with the images
embedded from docs/article-img/ (no network needed).

    python3 scripts/article-docx.py [out.docx]     # default: 小黑盒文章.docx in the repo root (git-ignored)

Needs python-docx (pip install python-docx). Handles exactly the subset the article uses:
H1-H3, paragraphs, **bold**, [text](url) links, '> ' quotes, flat lists, image lines.
"""
import os, re, sys
from docx import Document
from docx.enum.dml import MSO_THEME_COLOR_INDEX
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "docs", "小黑盒文章.md")
IMG = os.path.join(ROOT, "docs", "article-img")
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "小黑盒文章.docx")


def link(par, text, url):
    rid = par.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), rid)
    r = OxmlElement("w:r"); rp = OxmlElement("w:rPr")
    c = OxmlElement("w:color"); c.set(qn("w:val"), "0563C1"); rp.append(c)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rp.append(u)
    r.append(rp); t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve"); r.append(t)
    h.append(r); par._p.append(h)


def inline(par, text, italic=False):
    for part in re.split(r"(\*\*[^*]+\*\*|\[[^\]]+\]\([^)]+\))", text):
        if not part: continue
        m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part)
        if m: link(par, m.group(1), m.group(2)); continue
        bold = part.startswith("**") and part.endswith("**")
        run = par.add_run(part[2:-2] if bold else part); run.bold = bold; run.italic = italic


def main():
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Microsoft YaHei"; st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    width = doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin
    for line in open(SRC, encoding="utf-8").read().split("\n"):
        if not line.strip(): continue
        m = re.match(r"^!\[[^\]]*\]\(\S+/([A-Za-z0-9._-]+)\)$", line)
        if m:
            doc.add_picture(os.path.join(IMG, m.group(1)), width=width); continue
        h = re.match(r"^(#{1,3}) (.*)$", line)
        if h: doc.add_heading(h.group(2), level=len(h.group(1))); continue
        if line.startswith("> "):
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.8); inline(p, line[2:], italic=True); continue
        if re.match(r"^[-*] ", line):
            inline(doc.add_paragraph(style="List Bullet"), line[2:]); continue
        n = re.match(r"^(\d+)\. (.*)$", line)
        if n:
            p = doc.add_paragraph(); inline(p, n.group(1) + ". " + n.group(2)); continue
        inline(doc.add_paragraph(), line)
    doc.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
